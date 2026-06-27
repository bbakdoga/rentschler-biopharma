#!/usr/bin/env python3
"""
Score a processed batch against the ground-truth 'Solution example' workbook.

Measures roughly what fraction of the known-correct values the OCR/extraction
recovered, so OCR changes (model, prompt, tuning) can be compared objectively.

    python scripts/score_vs_solution.py <batch_id> ["Solution example - Part 1.xlsx"]
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import re

import openpyxl
from db.session import get_session
from db.models import Field, Personnel, Signature, Batch


def _norm(s: str) -> str:
    return "".join(str(s).split()).lower()


_DATE_ISO = re.compile(r"(\d{4})-(\d{2})-(\d{2})")
_DATE_DE = re.compile(r"\b(\d{1,2})[.](\d{1,2})[.](\d{2,4})\b")


def _as_date(s: str) -> str | None:
    """Canonicalise a value to YYYY-MM-DD if it looks like a date, else None."""
    s = str(s)
    m = _DATE_ISO.search(s)
    if m:
        return f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
    m = _DATE_DE.search(s)
    if m:
        d, mo, y = m.groups()
        if len(y) == 2:
            y = "20" + y
        return f"{y}-{int(mo):02d}-{int(d):02d}"
    return None


def ground_truth(path: str):
    """(parameter, value) rows from the solution; column F = parameter,
    column J = value. Checkbox states are reported separately."""
    ws = openpyxl.load_workbook(path, data_only=True)["Sheet1"]
    values, checks = [], 0
    for r in range(2, ws.max_row + 1):
        p, v = ws.cell(r, 6).value, ws.cell(r, 10).value
        if not p or v is None:
            continue
        sv = str(v).strip()
        if sv in ("Checked", "Not checked"):
            checks += 1
        else:
            values.append((str(p).strip(), sv))
    return values, checks


def main() -> int:
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    batch_id = int(sys.argv[1])
    sol = sys.argv[2] if len(sys.argv) > 2 else "Solution example - Part 1.xlsx"

    gt, n_checks = ground_truth(sol)

    # Gather everything the tool captured for this batch, across all tables.
    s = get_session()
    raw = []
    for f in s.query(Field).filter(Field.batch_id == batch_id).all():
        raw += [f.parsed_value, f.raw_value]
    for p in s.query(Personnel).filter(Personnel.batch_id == batch_id).all():
        raw += [p.name, p.kuerzel]
    for sig in s.query(Signature).filter(Signature.batch_id == batch_id).all():
        raw += [sig.kuerzel, sig.date]
    b = s.get(Batch, batch_id)
    if b:
        raw += [b.doc_nr, b.batch_no, b.production_code, b.production_code_ds,
                b.sap_material_nr, b.man_nr, b.process_step]
    s.close()

    raw = [str(x) for x in raw if x]
    tool = {_norm(x) for x in raw}
    tool_dates = {d for d in (_as_date(x) for x in raw) if d}

    hit, miss = 0, []
    for p, v in gt:
        gd = _as_date(v)
        if gd is not None:
            ok = gd in tool_dates
        else:
            nv = _norm(v)
            ok = bool(nv) and any(nv in tv for tv in tool)
        if ok:
            hit += 1
        else:
            miss.append((p, v))

    pct = hit / len(gt) * 100 if gt else 0
    print(f"Batch {batch_id} vs {sol}")
    print(f"  value fields matched: {hit}/{len(gt)} = {pct:.0f}%")
    print(f"  (ground-truth also has {n_checks} checkbox states, not scored here)")
    print("\n  sample of missed ground-truth values:")
    for p, v in miss[:25]:
        print(f"    {p[:36]:36} = {v[:32]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
