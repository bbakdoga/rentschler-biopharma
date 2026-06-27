from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from db.models import Batch, ValidationResult, Field, Signature, Personnel
from db.session import get_session
from config import OCR_CONFIDENCE_WARN

_H_FILL  = PatternFill("solid", fgColor="1F3864")
_H_FONT  = Font(color="FFFFFF", bold=True)
_E_FILL  = PatternFill("solid", fgColor="FFCCCC")
_W_FILL  = PatternFill("solid", fgColor="FFF2CC")
_I_FILL  = PatternFill("solid", fgColor="DDEBF7")
_G_FILL  = PatternFill("solid", fgColor="E2EFDA")
_THIN    = Side(style='thin', color='BBBBBB')
_BORDER  = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _header_row(ws, headers: list[str], row: int = 1):
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.fill   = _H_FILL
        c.font   = _H_FONT
        c.border = _BORDER
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def _autofit(ws, max_width: int = 60):
    for col in ws.columns:
        best = max((len(str(cell.value or '')) for cell in col), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(best + 2, max_width)


class ExcelExporter:

    def export(self, batch_id: int, out_path: str):
        session = get_session()
        try:
            batch = session.get(Batch, batch_id)
            if not batch:
                raise ValueError(f"Batch {batch_id} not found")

            wb = openpyxl.Workbook()
            self._summary(wb, batch, session)
            self._validation(wb, batch_id, session)
            self._signatures(wb, batch_id, session)
            self._fields(wb, batch_id, session)
            self._personnel(wb, batch_id, session)

            if 'Sheet' in wb.sheetnames:
                del wb['Sheet']

            wb.save(out_path)
        finally:
            session.close()

    # ── sheets ────────────────────────────────────────────────────────────────
    def _summary(self, wb, batch: Batch, session):
        ws = wb.create_sheet("Summary", 0)
        ws.sheet_view.showGridLines = False

        ws.merge_cells('A1:E1')
        ws['A1'] = "BPR Validation Report"
        ws['A1'].font = Font(size=18, bold=True, color="1F3864")

        info = [
            ("Document Nr",       batch.doc_nr),
            ("Batch No.",         batch.batch_no),
            ("Production Code",   batch.production_code),
            ("Process Step",      batch.process_step),
            ("SAP Material Nr",   batch.sap_material_nr),
            ("MAN Nr",            batch.man_nr),
            ("Revision",          batch.revision),
            ("GQQ generated",     batch.gqq_generated_at),
            ("File",              str(Path(batch.file_path or '').name)),
            ("Status",            batch.status),
            ("Processed at",      batch.processed_at.strftime('%d.%m.%Y %H:%M')
                                   if batch.processed_at else '—'),
            ("Report generated",  datetime.now().strftime('%d.%m.%Y %H:%M')),
        ]
        for i, (label, value) in enumerate(info, 3):
            ws.cell(row=i, column=1, value=label).font = Font(bold=True)
            ws.cell(row=i, column=2, value=value or '—')

        # counts
        from sqlalchemy import func
        counts = dict(
            session.query(ValidationResult.severity,
                          func.count(ValidationResult.id))
            .filter(ValidationResult.batch_id == batch.id)
            .group_by(ValidationResult.severity).all()
        )
        errors   = counts.get('error', 0)
        warnings = counts.get('warning', 0)
        infos    = counts.get('info', 0)

        row = len(info) + 5
        ws.merge_cells(f'A{row}:E{row}')
        ws.cell(row=row, column=1, value="Validation Summary").font = Font(bold=True, size=12)
        row += 1
        for label, val, fill in [
            ("Errors",   errors,   _E_FILL),
            ("Warnings", warnings, _W_FILL),
            ("Info",     infos,    _I_FILL),
        ]:
            c1 = ws.cell(row=row, column=1, value=label); c1.font = Font(bold=True)
            c2 = ws.cell(row=row, column=2, value=val);   c2.fill = fill if val else PatternFill()
            row += 1

        row += 1
        ws.cell(row=row, column=1, value="OVERALL STATUS").font = Font(bold=True, size=12)
        status_cell = ws.cell(row=row, column=2,
                              value="PASS" if errors == 0 else "FAIL")
        status_cell.font = Font(bold=True, size=12,
                                color="00AA00" if errors == 0 else "FF0000")

        ws.column_dimensions['A'].width = 22
        ws.column_dimensions['B'].width = 28

    def _validation(self, wb, batch_id, session):
        ws = wb.create_sheet("Validation Results")
        headers = ['#', 'Severity', 'Rule ID', 'Rule Name', 'Section', 'Page', 'Message', 'Field']
        _header_row(ws, headers)
        ws.row_dimensions[1].height = 22

        results = (
            session.query(ValidationResult)
            .filter(ValidationResult.batch_id == batch_id)
            .order_by(ValidationResult.severity, ValidationResult.rule_id)
            .all()
        )
        sev_fill = {'error': _E_FILL, 'warning': _W_FILL, 'info': _I_FILL}

        for i, r in enumerate(results, 1):
            fill = sev_fill.get(r.severity or '', PatternFill())
            row  = i + 1
            values = [i, (r.severity or '').upper(), r.rule_id,
                      r.rule_name, r.section, r.page_num, r.message, r.field_name]
            for col, v in enumerate(values, 1):
                c = ws.cell(row=row, column=col, value=v)
                c.fill   = fill
                c.border = _BORDER
                c.alignment = Alignment(wrap_text=(col == 7))

        _autofit(ws)

    def _signatures(self, wb, batch_id, session):
        ws = wb.create_sheet("Signatures")
        _header_row(ws, ['Role', 'Section', 'Kürzel', 'Date', 'Page'])
        sigs = session.query(Signature).filter(Signature.batch_id == batch_id).all()
        for i, s in enumerate(sigs, 2):
            for col, v in enumerate([s.role, s.section, s.kuerzel, s.date, s.page_num], 1):
                ws.cell(row=i, column=col, value=v).border = _BORDER
        _autofit(ws)

    def _fields(self, wb, batch_id, session):
        ws = wb.create_sheet("Extracted Fields")
        _header_row(ws, ['Section', 'Field Name', 'Raw Value', 'Parsed Value',
                         'Unit', 'Page', 'OCR Conf'])
        fields = session.query(Field).filter(Field.batch_id == batch_id).all()
        for i, f in enumerate(fields, 2):
            page_num = f.page.page_num if f.page else ''
            conf = f.confidence
            ctext = '' if conf is None else f'{int(conf)}%'
            uncertain = conf is not None and conf < OCR_CONFIDENCE_WARN
            for col, v in enumerate([f.section, f.field_name, f.raw_value,
                                      f.parsed_value, f.unit, page_num, ctext], 1):
                cell = ws.cell(row=i, column=col, value=v)
                cell.border = _BORDER
                if uncertain:                       # highlight shaky reads
                    cell.fill = _W_FILL
        _autofit(ws)

    def _personnel(self, wb, batch_id, session):
        ws = wb.create_sheet("Personnel")
        _header_row(ws, ['Name', 'Kürzel'])
        people = session.query(Personnel).filter(Personnel.batch_id == batch_id).all()
        for i, p in enumerate(people, 2):
            ws.cell(row=i, column=1, value=p.name).border   = _BORDER
            ws.cell(row=i, column=2, value=p.kuerzel).border = _BORDER
        _autofit(ws)
